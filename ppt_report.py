"""Power BI-style 16:9 PowerPoint deck (same charts as the PDF)."""
from __future__ import annotations

import io
from datetime import date

import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

from pdf_report import _short, chart_slides, kpi_items


def _rgb(hex_colour):
    h = hex_colour.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


BLUE, RED, NO_GREY = _rgb("#118DFF"), _rgb("#D64550"), _rgb("#E6E6E6")
DARK, CANVAS, INK, MUTED = _rgb("#252423"), _rgb("#F3F2F1"), _rgb("#252423"), _rgb("#605E5C")
LIGHT, WHITE, SHADOW = _rgb("#C8C6C4"), _rgb("#FFFFFF"), _rgb("#E1DFDD")
NAVY, HEADER_SUBTEXT = _rgb("#12239E"), _rgb("#E3F1FF")
FONT = "Segoe UI"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
HEADER_H = Inches(0.95)


def _rect(slide, x, y, w, h, colour, rounded=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if rounded:
        shape.adjustments[0] = 0.04
    shape.fill.solid()
    shape.fill.fore_color.rgb = colour
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _text(slide, text, x, y, w, h, size=14, bold=False, colour=INK):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size, p.font.bold, p.font.name = Pt(size), bold, FONT
        p.font.color.rgb = colour
    return box


def _card(slide, x, y, w, h, top=None):
    _rect(slide, x + Emu(12000), y + Emu(15000), w, h, SHADOW, rounded=True)
    _rect(slide, x, y, w, h, WHITE, rounded=True)
    if top is not None:
        _rect(slide, x + Inches(0.05), y, w - Inches(0.1), Inches(0.06), top)


def _base_slide(prs, title, subtitle, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = CANVAS
    _rect(slide, 0, 0, SLIDE_W, HEADER_H, BLUE)                 # sky-blue header band
    _rect(slide, 0, 0, Inches(0.09), HEADER_H, NAVY)
    _text(slide, title, Inches(0.4), Inches(0.14), Inches(12), Inches(0.5), size=24, bold=True, colour=WHITE)
    if subtitle:
        _text(slide, subtitle, Inches(0.4), Inches(0.58), Inches(12), Inches(0.3), size=12, colour=HEADER_SUBTEXT)
    _text(slide, footer, Inches(0.4), Inches(7.15), Inches(11), Inches(0.25), size=9, colour=MUTED)
    return slide


def _figure_card(slide, fig):
    x, y = Inches(0.35), HEADER_H + Inches(0.2)
    w, h = SLIDE_W - Inches(0.7), SLIDE_H - y - Inches(0.45)
    _card(slide, x, y, w, h)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight", facecolor="white")
    w_in, h_in = fig.get_size_inches()
    plt.close(fig)
    buf.seek(0)
    pad = Inches(0.2)
    scale = min((w - 2 * pad) / w_in, (h - 2 * pad) / h_in)
    pw, ph = Emu(int(w_in * scale)), Emu(int(h_in * scale))
    slide.shapes.add_picture(buf, int(x + (w - pw) / 2), int(y + pad), pw, ph)


def build_pptx(
    df: pd.DataFrame,
    report_title: str,
    summary: pd.DataFrame,
    cat_cols: list,
    filters_text: str,
    group: str | None = None,
    chosen: list | None = None,
) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    footer = f"{report_title} · {date.today():%d %b %Y} · Aggregate results only, no individual data"

    # ---- Title slide ----
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BLUE                # sky-blue cover
    _rect(slide, 0, 0, Inches(0.2), SLIDE_H, NAVY)
    _rect(slide, Inches(0.9), Inches(2.25), Inches(1.1), Inches(0.07), WHITE)
    _text(slide, report_title, Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.0), size=40, bold=True, colour=WHITE)
    _text(slide, f"{len(df):,} responses", Inches(0.9), Inches(3.55), Inches(11.5), Inches(0.5), size=20, colour=HEADER_SUBTEXT)
    _text(slide, filters_text, Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.0), size=14, colour=HEADER_SUBTEXT)
    _text(slide, f"Generated {date.today():%d %B %Y}", Inches(0.9), Inches(6.7), Inches(8), Inches(0.4),
          size=12, colour=HEADER_SUBTEXT)

    # ---- Summary slide: KPI cards + highest / lowest ----
    slide = _base_slide(prs, "Summary", "Key indicators and the questions with the highest and lowest % Yes", footer)
    items = kpi_items(df, summary)
    gap = Inches(0.2)
    left = Inches(0.35)
    kw = int((SLIDE_W - 2 * left - gap * (len(items) - 1)) / len(items))
    ky, kh = HEADER_H + Inches(0.2), Inches(1.45)
    for i, (value, label, red) in enumerate(items):
        x = left + i * (kw + gap)
        _card(slide, x, ky, kw, kh, top=RED if red else BLUE)
        _text(slide, value, x + Inches(0.2), ky + Inches(0.18), kw - Inches(0.3), Inches(0.6), size=30, bold=True,
              colour=RED if red else INK)
        _text(slide, label, x + Inches(0.2), ky + Inches(0.8), kw - Inches(0.35), Inches(0.6), size=10.5, colour=MUTED)

    top = summary.sort_values("% Yes", ascending=False)
    cy = ky + kh + Inches(0.2)
    ch = SLIDE_H - cy - Inches(0.45)
    cw = int((SLIDE_W - 2 * left - gap) / 2)
    for col, (label, rows) in enumerate([("Highest % Yes", top.head(6)), ("Lowest % Yes", top.tail(6).iloc[::-1])]):
        x = left + col * (cw + gap)
        _card(slide, x, cy, cw, ch)
        _text(slide, label, x + Inches(0.25), cy + Inches(0.15), cw, Inches(0.35), size=14, bold=True)
        row_h = int((ch - Inches(0.6)) / 6)
        for i, (_, r) in enumerate(rows.iterrows()):
            ry = cy + Inches(0.6) + i * row_h
            red = bool(r.get("Highlight", False))
            colour = RED if red else BLUE
            _text(slide, f"{r['% Yes']:.0f}%", x + Inches(0.25), ry, Inches(0.9), Inches(0.35), size=15, bold=True,
                  colour=RED if red else INK)
            bar_w = Inches(0.9)
            _rect(slide, x + Inches(0.25), ry + Inches(0.36), bar_w, Inches(0.07), NO_GREY)
            if r["% Yes"] > 0:
                _rect(slide, x + Inches(0.25), ry + Inches(0.36), max(1, int(bar_w * r["% Yes"] / 100)),
                      Inches(0.07), colour)
            q = _short(r["Question"], 1000).replace("\n", " ")
            _text(slide, q if len(q) < 120 else q[:118] + "…", x + Inches(1.35), ry + Inches(0.03),
                  cw - Inches(1.6), row_h, size=10.5)

    # ---- Chart slides (same as the PDF) ----
    for title, subtitle, fig in chart_slides(df, summary, cat_cols, group, chosen):
        slide = _base_slide(prs, title, subtitle, footer)
        _figure_card(slide, fig)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
