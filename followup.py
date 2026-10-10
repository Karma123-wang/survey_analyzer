"""CONFIDENTIAL follow-up list for the school counsellor.

Lists students who answered 'Yes' to the red (key) indicators, with name, class, section, age and their written
reason, so the counsellor can follow up with each student privately.
This list is NEVER included in the PDF, PowerPoint or Interactive report.
"""
from __future__ import annotations

import io
from datetime import date

import pandas as pd

TEXT_START = ("reason", "specify", "please specify", "write the name")


def _find(cols, words, exclude=()):
    for c in cols:
        low = c.lower().strip()
        if any(w in low for w in words) and not any(x in low for x in exclude):
            return c
    return None


def identity_columns(df: pd.DataFrame) -> dict:
    cols = list(df.columns)
    return {
        "Name": next((c for c in cols if c.lower().strip(" :") == "name"), None)
                or _find(cols, ["student name", "name of student", "full name"]),
        "Class": _find(cols, ["class", "grade"]),
        "Section": _find(cols, ["section"]),
        "Age": next((c for c in cols if c.lower().strip(" :") == "age"), None),
        "Sex": _find(cols, ["sex", "gender"]),
    }


def _reason_col(df: pd.DataFrame, question: str):
    """The written-reason column that follows a Yes/No question in the form, if there is one."""
    cols = list(df.columns)
    i = cols.index(question)
    if i + 1 < len(cols) and cols[i + 1].lower().startswith(TEXT_START):
        return cols[i + 1]
    return None


def _clean_reason(v) -> str:
    s = "" if pd.isna(v) else str(v).strip()
    return "" if s.strip("-_*=.() ").lower() in ("", "no", "nil", "none", "na", "n/a") else s


def followup_table(df: pd.DataFrame, red_questions: dict) -> pd.DataFrame:
    """red_questions = {original column: short label}. One row per student with at least one red 'Yes'."""
    ids = identity_columns(df)
    rows = []
    for idx, rec in df.iterrows():
        flags, reasons = [], []
        for q, label in red_questions.items():
            if q in df.columns and rec.get(q) == "Yes":
                flags.append(label)
                rc = _reason_col(df, q)
                rtxt = _clean_reason(rec.get(rc)) if rc else ""
                if rtxt:
                    reasons.append(f"{label}: {rtxt}")
        if flags:
            row = {k: ("" if c is None or pd.isna(rec.get(c)) else str(rec.get(c)).strip()) for k, c in ids.items()}
            row["Number of red indicators"] = len(flags)
            row["Answered Yes to"] = "\n".join(f"• {f}" for f in flags)
            row["Student's written reason"] = "\n".join(reasons)
            rows.append(row)
    out = pd.DataFrame(rows)
    if out.empty:
        return out
    from pdf_report import _natural_key
    out["_c"] = out["Class"].map(_natural_key)
    out["_a"] = out["Age"].map(_natural_key)
    out = out.sort_values(["Number of red indicators", "_c", "Section", "_a", "Name"],
                          ascending=[False, True, True, True, True]).drop(columns=["_c", "_a"])
    out.insert(0, "No.", range(1, len(out) + 1))
    out["Follow-up done (date / remarks)"] = ""
    return out.reset_index(drop=True)


def build_followup_xlsx(table: pd.DataFrame, school: str, counsellor: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Follow-up list"
    red = "C0392B"
    ws["A1"] = "CONFIDENTIAL – FOR THE SCHOOL COUNSELLOR ONLY"
    ws["A1"].font = Font(bold=True, size=14, color=red)
    ws["A2"] = (f"{school} · Students who answered Yes to the key (red) indicators · prepared for {counsellor or 'the counsellor'}"
                f" · {date.today():%d %b %Y}")
    ws["A3"] = ("Do not share, print for meetings, or attach to reports. Use only for private, supportive follow-up. "
                "Store securely and delete when follow-up is complete.")
    ws["A2"].font = Font(italic=True, color="605E5C")
    ws["A3"].font = Font(bold=True, color=red)
    thin = Side(style="thin", color="D0D7E2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = list(table.columns)
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=5, column=j, value=h)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=red), border
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, rec in enumerate(table.itertuples(index=False), start=6):
        for j, v in enumerate(rec, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = border
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if table.iloc[i - 6]["Number of red indicators"] > 1:
            for j in range(1, len(headers) + 1):
                ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor="FBE3E5")
    widths = {"No.": 5, "Name": 24, "Class": 8, "Section": 8, "Age": 6, "Sex": 8, "Number of red indicators": 10,
              "Answered Yes to": 42, "Student's written reason": 55, "Follow-up done (date / remarks)": 30}
    for j, h in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(row=5, column=j).column_letter].width = widths.get(h, 14)
    ws.freeze_panes = "C6"
    ws.row_dimensions[5].height = 32
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = "5:5"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
